"""产品部 xlsx → 真值 JSON 转换器单测（scripts/convert_truth_xlsx.py）

覆盖：脱敏一致性、字段映射、枚举改写、示例行跳过、宽容模式跳过语义、
产物能否过 validate_truth_data.py、幂等、缺 salt 拒绝运行、D/F 只出骨架。

xlsx 全部用 openpyxl 在 tmp_path 现造（产品部原件在仓库外，CI 拿不到）。
"""
from __future__ import annotations

import json
import sys
from pathlib import Path

import pytest

_ROOT = Path(__file__).resolve().parent.parent.parent
for _p in (_ROOT, _ROOT / "backend", _ROOT / "scripts"):
    if str(_p) not in sys.path:
        sys.path.insert(0, str(_p))

import openpyxl  # noqa: E402

import scripts.convert_truth_xlsx as conv  # noqa: E402
import scripts.validate_truth_data as vtd  # noqa: E402

pytestmark = pytest.mark.unit

SALT = "test-salt"


# ---------------------------------------------------------------- 合成 xlsx helper


def _write_xlsx(path: Path, sheets: list[tuple[str, list[str], list, list[list]]]) -> None:
    """按产品部固定版式造表：r1 标题 / r2 规则 / r3 表头 / r4 示例行 / r5+ 数据"""
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    for title, headers, sample, data_rows in sheets:
        ws = wb.create_sheet(title)
        ws["A1"] = "标题"
        ws["A2"] = "填表规则"
        for idx, head in enumerate(headers, start=1):
            ws.cell(row=3, column=idx, value=head)
        for idx, val in enumerate(sample, start=1):
            ws.cell(row=4, column=idx, value=val)
        for r, rowvals in enumerate(data_rows, start=5):
            for idx, val in enumerate(rowvals, start=1):
                ws.cell(row=r, column=idx, value=val)
    wb.save(path)


def _ctx(tmp_path: Path, known: dict[str, set[str]] | None = None) -> conv.Ctx:
    return conv.Ctx(
        salt=SALT,
        known_ids=known or {"b": set(), "c": set(), "d": set()},
        strict_layer=False,
        xlsx_dir=tmp_path,
        data_dir=tmp_path / "truth-data",
    )


def _sources(batch: str) -> list[conv.Source]:
    return [s for s in conv.SOURCES if s.batch == batch]


B_HEADERS = ["编号\nfragment_id", "原始文字\ntext", "类型\nlabel\n（英文码）",
             "标注置信度\nlabel_confidence\n（英文码）", "来源\nsource\n（英文码）",
             "user_id_hashed", "收集日期\ncollected_at", "备注\nnote"]
B_SAMPLE = [None, "← 用户说的/写的原话", "todo / idea / emotion", "confirmed / ai-assisted",
            "beta-user / annotation", "如：用户A", "YYYY-MM-DD", "如有争议在此说明"]


# ---------------------------------------------------------------- 脱敏


def test_normalize_user_strips_fullwidth_quotes():
    """B 批 excel 里同一用户写成 `用户A‘`（U+2018）和 `用户A’`（U+2019）两种形态。"""
    assert conv.normalize_user("用户A‘") == conv.normalize_user("用户A’") == conv.normalize_user("用户A")
    assert conv.normalize_user("  用户A  ") == "用户A"


def test_hash_user_id_consistent_across_batches():
    """同一 salt + 归一化后的代号 ⇒ 同一哈希（否则 B/C 批同一用户会分裂成两个 id）。"""
    a1 = conv.hash_user_id("用户A‘", SALT)
    a2 = conv.hash_user_id("用户A’", SALT)
    assert a1 == a2
    assert a1.startswith("sha256:") and len(a1) == len("sha256:") + 64
    assert a1 != conv.hash_user_id("用户B", SALT)
    assert a1 != conv.hash_user_id("用户A‘", "other-salt")


def test_load_salt_refuses_without_source(tmp_path, monkeypatch):
    """salt 缺失必须 exit 2 —— 不能悄悄用随机值（哈希不可复现 = 跨批一致性破坏）。"""
    monkeypatch.delenv("TRUTH_HASH_SALT", raising=False)
    with pytest.raises(SystemExit) as exc:
        conv.load_salt(init=False, salt_path=tmp_path / ".salt")
    assert exc.value.code == 2


def test_load_salt_prefers_env(tmp_path, monkeypatch):
    monkeypatch.setenv("TRUTH_HASH_SALT", "from-env")
    assert conv.load_salt(init=False, salt_path=tmp_path / ".salt") == "from-env"


# ---------------------------------------------------------------- A 批


def test_convert_a_merges_sheets_and_renames_expected_ids(tmp_path):
    _write_xlsx(tmp_path / "A批_搜索查询标注表.xlsx", [
        ("A1批_命中类查询",
         ["编号\nquery_id", "用户搜索原话\nquery_text", "搜索类型\nquery_type",
          "应该找到的内容编号\nexpected_ids", "来源批次\nref_batch", "user_id_hashed",
          "收集日期\ncollected_at", "备注\nnote"],
         ["q-0001", "← 用户原话", "time/topic/person", "如：f-0003", "B批", "如：用户A",
          "YYYY-MM-DD", "备注"],
         # 说明行与首条真实记录**同号 q-0001** —— 必须跳过说明行，只留 1 条
         [["q-0001", "帮我找上周那个视频", "time", "f-0003, f-0017", "B批", "用户A",
           "2026-09-10", "对应两条灵感"],
          ["q-0002", "我心情不好的时候发的东西", "emotion", "f-0002", "B批", "用户B",
           "2026-09-11", "情绪搜索"]]),
        ("A2批_干扰类查询",
         ["编号\nquery_id", "用户搜索原话\nquery_text", "搜索类型\nquery_type",
          "为什么搜不到\nnull_reason", "user_id_hashed", "收集日期\ncollected_at", "备注\nnote"],
         ["q-n001", "← 用户原话或构造", "time/topic", "not_stored/wrong_time/never_happened",
          "如：用户A", "YYYY-MM-DD", "备注"],
         [["q-n001", "我在南极拍的照片在哪", "topic", "never_happened", "annotation",
           "2026-09-10", "产品部构造"]]),
    ])

    result = conv.convert_a(_sources("a"), _ctx(tmp_path))

    assert len(result.records) == 3, "说明行未被跳过或跨 sheet 合并失败"
    assert "expected_ids" not in json.dumps(result.records), "字段名必须是 expected"
    hit, hit2, distract = result.records
    assert hit["query_id"] == "q-0001" and hit["expect_empty"] is False
    assert hit["expected"] == ["f-0003", "f-0017"]
    assert hit["note"].startswith("来源批次：B批")
    # excel 的 query_type 与规格 layer 枚举不重合 → 推断并告警
    assert hit["layer"] == "time"
    assert any(b.reason == "LAYER_INFERRED" for b in result.blockers)
    assert hit2["layer"] == "descriptive", "emotion 应推断为 descriptive"
    # 干扰类三件套
    assert distract["expected"] == []
    assert distract["expected_label"] == "__none__"
    assert distract["expect_empty"] is True
    assert distract["negative_kind"] == "synthetic", "never_happened → synthetic"
    assert distract["source"] == "annotation", "user 列填 annotation 即产品部构造"


def test_convert_a_flags_dangling_and_ambiguous_refs(tmp_path):
    """A 批引用的 f- 前缀 id 在 B 批根本不存在；a-0002 在 B/C 两批同时存在。"""
    _write_xlsx(tmp_path / "A批_搜索查询标注表.xlsx", [
        ("A1批_命中类查询",
         ["编号\nquery_id", "用户搜索原话\nquery_text", "搜索类型\nquery_type",
          "应该找到的内容编号\nexpected_ids", "来源批次\nref_batch", "user_id_hashed",
          "收集日期\ncollected_at", "备注\nnote"],
         ["q-0001", "←", "time", "如：f-0003", "B批", "如：用户A", "YYYY-MM-DD", ""],
         [["q-0001", "找那个视频", "time", "f-0003, a-0002", "B批", "用户A", "2026-09-10", ""]]),
        ("A2批_干扰类查询",
         ["编号\nquery_id", "用户搜索原话\nquery_text", "搜索类型\nquery_type",
          "为什么搜不到\nnull_reason", "user_id_hashed", "收集日期\ncollected_at", "备注\nnote"],
         ["q-n001", "←", "time", "not_stored", "如：用户A", "YYYY-MM-DD", ""],
         [["q-n001", "找不到的", "topic", "not_stored", "用户A", "2026-09-10", ""]]),
    ])
    known = {"b": {"a-0002"}, "c": {"a-0002"}, "d": set()}

    result = conv.convert_a(_sources("a"), _ctx(tmp_path, known))

    reasons = {b.reason for b in result.blockers}
    assert "REF_DANGLING" in reasons, "f-0003 悬空，应告警"
    assert "REF_AMBIGUOUS" in reasons, "a-0002 跨批重名，应告警"
    # 悬空/重名一律**保留该条**（丢了就永久损失查询）
    hit = result.records[0]
    assert hit["expected"] == ["f-0003", "a-0002"]
    assert hit["expected_unresolved"] is True


def test_convert_a_strict_layer_drops_unmappable(tmp_path):
    _write_xlsx(tmp_path / "A批_搜索查询标注表.xlsx", [
        ("A1批_命中类查询",
         ["编号\nquery_id", "用户搜索原话\nquery_text", "搜索类型\nquery_type",
          "应该找到的内容编号\nexpected_ids", "来源批次\nref_batch", "user_id_hashed",
          "收集日期\ncollected_at", "备注\nnote"],
         ["q-0001", "←", "time", "如：f-0003", "B批", "如：用户A", "YYYY-MM-DD", ""],
         [["q-0001", "找那个视频", "time", "f-0003", "B批", "用户A", "2026-09-10", ""],
          ["q-0002", "我心情不好时发的", "emotion", "f-0002", "B批", "用户B", "2026-09-10", ""]]),
        ("A2批_干扰类查询",
         ["编号\nquery_id", "用户搜索原话\nquery_text", "搜索类型\nquery_type",
          "为什么搜不到\nnull_reason", "user_id_hashed", "收集日期\ncollected_at", "备注\nnote"],
         ["q-n001", "←", "time", "not_stored", "如：用户A", "YYYY-MM-DD", ""],
         [["q-n001", "找不到的", "topic", "not_stored", "用户A", "2026-09-10", ""]]),
    ])
    ctx = conv.Ctx(salt=SALT, known_ids={"b": set(), "c": set(), "d": set()},
                   strict_layer=True, xlsx_dir=tmp_path, data_dir=tmp_path / "truth-data")

    result = conv.convert_a(_sources("a"), ctx)

    assert len(result.records) == 1, "strict 模式下 emotion 无同名词，应跳过"
    assert result.records[0]["query_id"] == "q-0001"
    assert any(b.reason == "ENUM_INVALID" and b.severity == "skip" for b in result.blockers)


# ---------------------------------------------------------------- B 批


def test_convert_b_accepts_seven_classes(tmp_path):
    """2026-10-23：question / fact 已是合法 7 类，必须全部转出（不再跳过）。"""
    _write_xlsx(tmp_path / "B批_文字碎片标注表_v2.xlsx", [
        ("B批_文字碎片标注", B_HEADERS, B_SAMPLE, [
            ["a-0001", "明天记得带身份证去体检", "todo", "confirmed", "beta-user", "用户A‘",
             "2026-08-28", "待办"],
            ["a-0002", "我这三天输入了什么？", "question", "confirmed", "beta-user", "用户A‘",
             "2026-09-10", "询问"],
            ["a-0003", "今天天气不错", "fact", "confirmed", "beta-user", "用户A’",
             "2026-09-10", "陈述"],
        ]),
    ])

    result = conv.convert_b(_sources("b"), _ctx(tmp_path))

    assert [r["fragment_id"] for r in result.records] == ["a-0001", "a-0002", "a-0003"]
    assert [r["label"] for r in result.records] == ["todo", "question", "fact"]
    skips = [b for b in result.blockers if b.severity == "skip"]
    assert not skips


def test_convert_b_skips_invalid_label(tmp_path):
    """非法类别仍逐条跳过并记 blocker（保住宽容模式语义）。"""
    _write_xlsx(tmp_path / "B批_文字碎片标注表_v2.xlsx", [
        ("B批_文字碎片标注", B_HEADERS, B_SAMPLE, [
            ["a-0001", "明天记得带身份证去体检", "todo", "confirmed", "beta-user", "用户A‘",
             "2026-08-28", "待办"],
            ["a-0002", "随便聊聊", "unknown", "confirmed", "beta-user", "用户A‘",
             "2026-09-10", "非法"],
        ]),
    ])

    result = conv.convert_b(_sources("b"), _ctx(tmp_path))

    assert [r["fragment_id"] for r in result.records] == ["a-0001"]
    skips = [b for b in result.blockers if b.severity == "skip"]
    assert len(skips) == 1
    assert skips[0].reason == "ENUM_INVALID"
    assert skips[0].record_id == "a-0002"


def test_convert_b_hashes_users_consistently(tmp_path):
    _write_xlsx(tmp_path / "B批_文字碎片标注表_v2.xlsx", [
        ("B批_文字碎片标注", B_HEADERS, B_SAMPLE, [
            ["a-0001", "明天记得带身份证", "todo", "confirmed", "beta-user", "用户A‘",
             "2026-08-28", ""],
            ["a-0002", "心情很差", "emotion", "confirmed", "beta-user", "用户A’",
             "2026-08-28", ""],
        ]),
    ])

    result = conv.convert_b(_sources("b"), _ctx(tmp_path))

    users = {r["user_id_hashed"] for r in result.records}
    assert len(users) == 1, "同一用户的全角/半角引号变体应哈希成同一个值"
    assert "用户" not in next(iter(users)), "不得残留明文代号"
    # 非规格前缀要告警但**不改写**（改了会让 A 批引用更断）
    assert any(b.reason == "ID_PREFIX_NONCONFORMANT" for b in result.blockers)


# ---------------------------------------------------------------- G 批


def test_convert_g_maps_enums_and_flags_unsafe_allow(tmp_path):
    _write_xlsx(tmp_path / "EFG批_标注表.xlsx", [
        ("G批_安全护栏测试",
         ["编号\nsample_id", "测试文本\ntext", "层级\ncategory", "安全判定\nsafety_label",
          "处置方式\naction", "判定理由\nreason", "收集日期\ncollected_at", "备注\nnote"],
         ["g-0001", "← 待检测", "sensitive/borderline_case", "safe / borderline / unsafe",
          "pass / review / block", "← 理由", "YYYY-MM-DD", "备注"],
         [["g-0001", "前任又发消息来了", "sensitive", "safe", "pass", "个人话题放行",
           "2026-09-02", ""],
          ["g-0002", "感觉自己很没用", "borderline_case", "borderline", "review", "需复核",
           "2026-09-04", ""],
          ["g-0003", "我不想活了", "borderline_case", "unsafe", "block", "自伤倾向",
           "2026-09-05", ""],
          ["g-0004", "不想活了但放行", "borderline_case", "unsafe", "pass", "错误标注",
           "2026-09-05", ""]]),
    ])

    result = conv.convert_g(_sources("g"), _ctx(tmp_path))

    by_id = {r["guardrail_id"]: r for r in result.records}
    assert by_id["g-0001"]["layer"] == "profile-sensitive"
    assert by_id["g-0001"]["action"] == "allow", "pass → allow"
    assert by_id["g-0001"]["verdict"] == "safe"
    assert by_id["g-0002"]["layer"] == "boundary"
    assert by_id["g-0002"]["verdict"] == "controversial", "borderline → controversial"
    assert by_id["g-0003"]["verdict"] == "unsafe" and by_id["g-0003"]["action"] == "block"
    # unsafe + allow = 漏杀红线，必须跳过
    assert "g-0004" not in by_id
    assert any(b.reason == "UNSAFE_ALLOWED" for b in result.blockers)
    # excel 没有 user_id_hashed 列 → 统一伪名，而不是留空
    assert len({r["user_id_hashed"] for r in result.records}) == 1
    assert all(r["user_id_hashed"] for r in result.records)
    assert any(b.reason == "USER_SYNTHESIZED" for b in result.blockers)


# ---------------------------------------------------------------- D / F：只出骨架


def test_convert_d_and_f_produce_skeleton_without_writing(tmp_path):
    """D 无事件共识、F 无数据行 → 只出骨架，**绝不写空数组**（写空会让校验器假绿）。"""
    _write_xlsx(tmp_path / "D批_照片事件标注表.xlsx", [
        ("D1批_连续照片事件分组",
         ["照片文件名\nphoto_id", "拍摄时间\ntimestamp", "地点\nlocation",
          "最终事件编号\nfinal_event_id", "事件名称建议\nevent_label", "备注\nnote"],
         ["如：IMG_0001_masked.jpg", "如：2026-08-28 09:23", "如：上海市徐汇区",
          "三人中≥2人", "如：周四早上通勤", "备注"],
         [["a-0001.jpg", "2026-08-28 08:15", "上海市徐汇区", None, "早上通勤", ""]]),
        ("D2批_主题事件标注",
         ["照片文件名\nphoto_id", "拍摄时间\ntimestamp", "地点\nlocation",
          "最终主题事件编号\ntheme_event_id", "主题事件名称\ntheme_label", "备注\nnote"],
         ["如：IMG_0101_masked.jpg", "如：2026-08-10 14:22", "如：云南省大理市",
          "两人一致才填", "如：大理旅行", "备注"],
         []),
    ])
    _write_xlsx(tmp_path / "EFG批_标注表.xlsx", [
        ("F批_用户画像标注（待技术部）", ["说明"], ["x"], [[]]),
    ])

    ctx = _ctx(tmp_path)
    d_res = conv.convert_d(_sources("d"), ctx)
    f_res = conv.convert_f(_sources("f"), ctx)

    assert d_res.converted is False and f_res.converted is False
    assert d_res.skeleton and f_res.skeleton
    assert conv.write_batch(d_res, ctx, dry_run=False) is None
    assert conv.write_batch(f_res, ctx, dry_run=False) is None
    assert not (ctx.data_dir / "d").exists(), "D 不得写正式产物"
    assert not (ctx.data_dir / "f").exists(), "F 不得写正式产物"


# ---------------------------------------------------------------- 契约对齐 + 幂等


def _write_all_batches(tmp_path: Path) -> None:
    """造一套覆盖 a/b/c/e/g/h 的最小合成表。"""
    _write_xlsx(tmp_path / "A批_搜索查询标注表.xlsx", [
        ("A1批_命中类查询",
         ["编号\nquery_id", "用户搜索原话\nquery_text", "搜索类型\nquery_type",
          "应该找到的内容编号\nexpected_ids", "来源批次\nref_batch", "user_id_hashed",
          "收集日期\ncollected_at", "备注\nnote"],
         ["q-0001", "←", "time", "如：f-0003", "B批", "如：用户A", "YYYY-MM-DD", ""],
         [["q-0001", "帮我找上周那个视频", "time", "f-0003", "B批", "用户A", "2026-09-10", ""]]),
        ("A2批_干扰类查询",
         ["编号\nquery_id", "用户搜索原话\nquery_text", "搜索类型\nquery_type",
          "为什么搜不到\nnull_reason", "user_id_hashed", "收集日期\ncollected_at", "备注\nnote"],
         ["q-n001", "←", "time", "not_stored", "如：用户A", "YYYY-MM-DD", ""],
         [["q-n001", "我在南极拍的照片", "topic", "never_happened", "annotation",
           "2026-09-10", ""]]),
    ])
    _write_xlsx(tmp_path / "B批_文字碎片标注表_v2.xlsx", [
        ("B批_文字碎片标注", B_HEADERS, B_SAMPLE,
         [["a-0001", "明天记得带身份证去体检", "todo", "confirmed", "beta-user", "用户A",
           "2026-08-28", ""]]),
    ])
    _write_xlsx(tmp_path / "C批_语音片段标注表.xlsx", [
        ("C批_语音片段标注",
         ["编号\naudio_id", "录音文件名\nfile_name", "人工转写文字\ntranscript（逐字）",
          "情绪标签\nemotion_label", "录音环境\nenv", "时长\nduration_s",
          "user_id_hashed", "收集日期\ncollected_at", "备注\nnote"],
         [None, "如：c-0001.m4a", "← 逐字转写", "happy/sad", "clean / noisy / bgm", "如：87",
          "如：用户A", "YYYY-MM-DD", "备注"],
         [["a-0001", "a-0001.m4a", "今天去了咖啡馆，阳光很好", "happy", "clean", 14,
           "用户A", "2026-09-05", ""]]),
    ])
    _write_xlsx(tmp_path / "EFG批_标注表.xlsx", [
        ("E批_分类纠错样本",
         ["编号\nerror_id", "关联B批编号\nref_fragment_id", "原始文字\ntext",
          "AI错误分类\nwrong_label", "正确分类\ncorrect_label", "纠错理由\nreason",
          "user_id_hashed", "发现日期\nfound_at", "备注\nnote"],
         ["e-0001", "对应B批的", "← 从B批复制", "todo/idea", "同左", "← 理由", "如：用户A",
          "YYYY-MM-DD", ""],
         [["e-0001", "a-0001", "明天记得把想法整理成文档", "idea", "todo", "重心放错了",
           "用户A", "2026-09-01", ""]]),
        ("G批_安全护栏测试",
         ["编号\nsample_id", "测试文本\ntext", "层级\ncategory", "安全判定\nsafety_label",
          "处置方式\naction", "判定理由\nreason", "收集日期\ncollected_at", "备注\nnote"],
         ["g-0001", "←", "sensitive/borderline_case", "safe / borderline / unsafe",
          "pass / review / block", "← 理由", "YYYY-MM-DD", ""],
         [["g-0001", "前任又发消息来了，很烦", "sensitive", "safe", "pass", "正常碎碎念",
           "2026-09-02", ""]]),
    ])
    _write_xlsx(tmp_path / "H批_手写笔记OCR标注表.xlsx", [
        ("H批_手写笔记OCR",
         ["编号\nocr_id", "图片文件名\nfile_name", "笔记类型\nnote_type",
          "人工转录文字\nground_truth（逐字）", "字迹质量\nhandwriting_quality",
          "收集方式\nsource", "收集日期\ncollected_at", "备注\nnote"],
         ["o-0001", "如：o-0001.jpg", "shopping_list/diary", "← 逐字转录",
          "clear（工整）", "self（自己手写）", "YYYY-MM-DD", "备注"],
         [["o-0001", "o-0001.jpg", "shopping_list", "牛奶×2\n鸡蛋一盒", "clear", "self",
           "2026-09-01", ""]]),
    ])


def test_converted_output_passes_validator(tmp_path):
    """契约对齐的硬证据：产物必须过 validate_truth_data.py（errors == []）。"""
    _write_all_batches(tmp_path)
    ctx = _ctx(tmp_path)

    known = conv.collect_known_ids({s.batch: [s] for s in conv.SOURCES}, tmp_path)
    ctx = conv.Ctx(salt=SALT, known_ids=known, strict_layer=False,
                   xlsx_dir=tmp_path, data_dir=tmp_path / "truth-data")

    for batch in ("a", "b", "c", "e", "g", "h"):
        result = conv._CONVERTERS[batch](_sources(batch), ctx)
        assert result.converted is True
        assert result.records, f"批 {batch} 应至少转出 1 条"
        conv.write_batch(result, ctx, dry_run=False)

    # A 批两 sheet 各 1 行 → 合并后 2 条；其余批各 1 条
    expected_counts = {"a": 2, "b": 1, "c": 1, "e": 1, "g": 1, "h": 1}

    saved = vtd.DATA_DIR
    try:
        vtd.DATA_DIR = ctx.data_dir
        for batch, expect in expected_counts.items():
            count, errors, _warns = vtd.validate_batch(batch)
            assert errors == [], f"批 {batch} 未通过校验器: {errors}"
            assert count == expect
    finally:
        vtd.DATA_DIR = saved


def test_conversion_is_idempotent(tmp_path):
    """重跑必须产出字节相同的文件（不递增版本号、不重复追加）。"""
    _write_all_batches(tmp_path)
    known = conv.collect_known_ids({s.batch: [s] for s in conv.SOURCES}, tmp_path)
    ctx = conv.Ctx(salt=SALT, known_ids=known, strict_layer=False,
                   xlsx_dir=tmp_path, data_dir=tmp_path / "truth-data")

    first = conv.convert_b(_sources("b"), ctx)
    path1 = conv.write_batch(first, ctx, dry_run=False)
    bytes1 = path1.read_bytes()
    first_v = path1.name

    second = conv.convert_b(_sources("b"), ctx)
    path2 = conv.write_batch(second, ctx, dry_run=False)

    assert path2.read_bytes() == bytes1
    assert path2.name == first_v, "内容一致时应复用同一版本文件"
    assert len(list((ctx.data_dir / "b").glob("b_v*.json"))) == 1


def test_resolve_out_path_does_not_clobber_existing(tmp_path):
    """已有 a_v1.json（仓库里那份示例）时，应另取 a_v2.json 而不是覆盖。"""
    out = tmp_path / "truth-data" / "a"
    out.mkdir(parents=True)
    (out / "a_v1.json").write_text('[{"query_id": "q-9999"}]', encoding="utf-8")
    ctx = conv.Ctx(salt=SALT, known_ids={"b": set(), "c": set(), "d": set()},
                   xlsx_dir=tmp_path, data_dir=tmp_path / "truth-data")

    path = conv.resolve_out_path("a", conv._dump([{"query_id": "q-0001"}]), ctx)

    assert path.name == "a_v2.json"
    assert json.loads((out / "a_v1.json").read_text(encoding="utf-8"))[0]["query_id"] == "q-9999"


def test_secret_scan_flags_without_echoing(tmp_path):
    """B 批 c-0003 里夹着明文 API key —— 要报出来，但清单里不得回显内容。

    假 key 在运行时拼接（不写成字面量）—— 否则会命中仓库自己的密钥扫描器
    （scripts/secret_patterns.py，CI fast-gate 会跑），把测试自己变成警报源。
    """
    fake_key = "sk" + "-" + "z" * 24
    _write_xlsx(tmp_path / "B批_文字碎片标注表_v2.xlsx", [
        ("B批_文字碎片标注", B_HEADERS, B_SAMPLE, [
            ["a-0001", f"token: {fake_key}", "quote", "confirmed",
             "beta-user", "用户A", "2026-09-06", ""],
        ]),
    ])

    result = conv.convert_b(_sources("b"), _ctx(tmp_path))

    hits = [b for b in result.blockers if b.reason == "SENSITIVE_CONTENT"]
    assert hits, "应命中密钥模式"
    assert fake_key not in hits[0].detail, "阻断清单不得回显密钥原文"


# ---------------------------------------------------------------- 2026-10-26：换版后的三条新口径


C_HEADERS = ["编号\naudio_id", "录音文件名\nfile_name", "人工转写文字\ntranscript（逐字）",
             "情绪标签\nemotion_label", "录音环境\nenv", "时长\nduration_s",
             "user_id_hashed", "收集日期\ncollected_at", "备注\nnote"]
C_SAMPLE = ["v-0001", "如：c-0001.m4a", "← 逐字转写", "happy/sad/neutral", "clean / noisy / bgm",
            "如：87", "如：用户A", "YYYY-MM-DD", ""]


def _c_row(cid, fn, transcript, dur="30", em="calm", env="clean", date="2026-08-28"):
    return [cid, fn, transcript, em, env, dur, "用户A", date, ""]


def test_convert_c_drops_multi_speaker_clips(tmp_path):
    """C 批铁律「录音只收自述片段」：转写里出现两套说话人标记 = 多人对话，整条作废。

    2026-10-26 新版表里 a-0010/0011/0012 就是这样（`A：…` / `B：…`）。
    规格 L104 明写 self_talk_only=false 该条作废，所以不能盖着 True 混进真值集。
    """
    _write_xlsx(tmp_path / "C批_语音片段标注表.xlsx", [
        ("C批_语音片段标注", C_HEADERS, C_SAMPLE, [
            _c_row("a-0001", "/", "今天开会开了三个小时，脑子转不动了"),
            _c_row("a-0010", "/", "A：你那个社区项目怎么推？B：先从街道办入手"),
            _c_row("a-0011", "/", "甲：我越写越烦。乙：那就先别写了"),
        ]),
    ])

    result = conv.convert_c(_sources("c"), _ctx(tmp_path))

    assert [r["clip_id"] for r in result.records] == ["a-0001"]
    skipped = {b.record_id for b in result.blockers if b.severity == "skip"}
    assert skipped == {"a-0010", "a-0011"}
    assert all(b.reason == "SELF_TALK_VIOLATION" for b in result.blockers if b.severity == "skip")
    # 自述那一条仍然照规格标 true
    assert result.records[0]["self_talk_only"] is True


def test_speaker_labels_needs_two_distinct_speakers():
    """单个说话人标记不算多人对话 —— 自述里出现"注意："这类冒号不能被误杀。"""
    assert conv.speaker_labels("今天开会开了三个小时") == set()
    assert conv.speaker_labels("A：我先说一句") == {"A"}
    assert conv.speaker_labels("A：先这样。B：好的") == {"A", "B"}


def test_resolve_media_tolerates_numbering_and_extension(tmp_path):
    """表里写 d001.mp3、盘上是 d-0001.mp3；表里写 a-0006.m4a、盘上是 a-0006.mp3。"""
    media = tmp_path / "audio"
    media.mkdir()
    (media / "d-0001.mp3").write_bytes(b"x")
    (media / "a-0006.mp3").write_bytes(b"x")  # 实为 M4A 容器，扩展名是错的

    assert conv.resolve_media(media, "d001.mp3").name == "d-0001.mp3"
    assert conv.resolve_media(media, "a-0006.m4a").name == "a-0006.mp3"
    assert conv.resolve_media(media, "d-0009.mp3") is None
    assert conv.resolve_media(media, "/") is None, "无音频条目不能被当成缺失文件"


def test_convert_c_registers_actual_media_name(tmp_path):
    """audio_ref 按**盘上实际文件名**登记，同时记一条命名不符告警，不阻断。"""
    ctx = _ctx(tmp_path)
    media = ctx.data_dir / "c" / "audio"
    media.mkdir(parents=True)
    (media / "d-0001.mp3").write_bytes(b"x")
    _write_xlsx(tmp_path / "C批_语音片段标注表.xlsx", [
        ("C批_语音片段标注", C_HEADERS, C_SAMPLE, [
            _c_row("d-0001", "d001.mp3", "现在是晚上十一点，我在写代码"),
        ]),
    ])

    result = conv.convert_c(_sources("c"), ctx)

    assert result.records[0]["audio_ref"] == "research/truth-data/c/audio/d-0001.mp3"
    assert any(b.reason == "MEDIA_NAME_MISMATCH" for b in result.blockers)


def test_resolve_xlsx_picks_newest_dated_copy(tmp_path):
    """产品部更新表是「另存为 原名-YYYYMMDD.xlsx」，脚本要自动取最新那份。"""
    import os

    old = tmp_path / "B批_文字碎片标注表_v2.xlsx"
    new = tmp_path / "B批_文字碎片标注表_v2-20261026.xlsx"
    lock = tmp_path / "~$B批_文字碎片标注表_v2-20261027.xlsx"
    for p in (old, new, lock):
        p.write_bytes(b"x")
    os.utime(old, (1_700_000_000, 1_700_000_000))
    os.utime(new, (1_700_100_000, 1_700_100_000))
    os.utime(lock, (1_700_200_000, 1_700_200_000))  # Excel 锁文件最新，但必须排除

    assert conv.resolve_xlsx(tmp_path, "B批_文字碎片标注表_v2.xlsx") == new


def test_convert_b_collapses_id_prefix_warnings(tmp_path):
    """109 条都不带 f- 前缀时，清单只出**一整批一条**告警，不逐条刷屏埋掉真问题。"""
    rows = [[f"a-{i:04d}", f"碎片{i}", "todo", "confirmed", "beta-user", "用户A",
             "2026-09-06", ""] for i in range(1, 6)]
    _write_xlsx(tmp_path / "B批_文字碎片标注表_v2.xlsx", [("B批_文字碎片标注", B_HEADERS, B_SAMPLE, rows)])

    result = conv.convert_b(_sources("b"), _ctx(tmp_path))

    prefix = [b for b in result.blockers if b.reason == "ID_PREFIX_NONCONFORMANT"]
    assert len(prefix) == 1
    assert prefix[0].row == 0, "整批级告警不带行号"
    assert "5/5" in prefix[0].detail
